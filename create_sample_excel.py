import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# Sheet 1: Laporan Penjualan
ws1 = wb.active
ws1.title = "Laporan Penjualan 2026"
ws1.views.sheetView[0].showGridLines = True

# Title
ws1.merge_cells("A1:G1")
title_cell = ws1["A1"]
title_cell.value = "PT NUSANTARA JAYA ABADI - LAPORAN REKAPITULASI PENJUALAN"
title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E293B")
title_cell.alignment = Alignment(horizontal="center", vertical="center")

ws1.merge_cells("A2:G2")
sub_cell = ws1["A2"]
sub_cell.value = "Periode: Q1 - Q3 Tahun Anggaran 2026 (Semua Divisi)"
sub_cell.font = Font(name="Calibri", size=10, italic=True, color="64748B")
sub_cell.alignment = Alignment(horizontal="center", vertical="center")

# Headers
headers = ["No", "Kode Transaksi", "Nama Produk / Layanan", "Kategori", "Qty", "Harga Satuan (Rp)", "Total Pendapatan (Rp)"]
header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

for col_idx, h in enumerate(headers, 1):
    cell = ws1.cell(row=4, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

# Data Rows
rows = [
    [1, "TRX-2026-001", "Paket Cloud Server Enterprise 64GB", "Infrastruktur", 15, 4500000, 67500000],
    [2, "TRX-2026-002", "Lisensi Software ERP Core v4.2", "Perangkat Lunak", 30, 2250000, 67500000],
    [3, "TRX-2026-003", "Jasa Integrasi API Gateway FinTech", "Layanan TI", 8, 12000000, 96000000],
    [4, "TRX-2026-004", "Maintenance & SLA Support Bulanan", "Pemeliharaan", 24, 3500000, 84000000],
    [5, "TRX-2026-005", "Implementasi Keamanan Siber & Firewall", "Keamanan", 5, 18500000, 92500000],
    [6, "TRX-2026-006", "Modul AI Analytics & Business Insights", "Kecerdasan Buatan", 12, 8500000, 102000000],
    [7, "TRX-2026-007", "Pelatihan Staf & Sertifikasi DevOps", "Pelatihan", 40, 1500000, 60000000],
    [8, "TRX-2026-008", "Dedicated Fiber Optic Backup Line", "Jaringan", 6, 7500000, 45000000],
]

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

for r_idx, row_data in enumerate(rows, 5):
    is_even = (r_idx % 2 == 0)
    for c_idx, val in enumerate(row_data, 1):
        cell = ws1.cell(row=r_idx, column=c_idx, value=val)
        cell.border = thin_border
        if is_even:
            cell.fill = zebra_fill
        if c_idx == 1:
            cell.alignment = Alignment(horizontal="center")
        elif c_idx == 2:
            cell.alignment = Alignment(horizontal="center")
            cell.font = Font(name="Calibri", bold=True)
        elif c_idx in [5, 6, 7]:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal="right")

# Total Row
tot_row = len(rows) + 5
ws1.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=6)
lbl_tot = ws1.cell(row=tot_row, column=1, value="TOTAL KESELURUHAN PENDAPATAN")
lbl_tot.font = Font(name="Calibri", size=11, bold=True, color="1E293B")
lbl_tot.alignment = Alignment(horizontal="right", vertical="center")
lbl_tot.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

val_tot = ws1.cell(row=tot_row, column=7, value=f"=SUM(G5:G{tot_row-1})")
val_tot.font = Font(name="Calibri", size=11, bold=True, color="059669")
val_tot.number_format = 'Rp #,##0'
val_tot.alignment = Alignment(horizontal="right", vertical="center")
val_tot.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

for c in range(1, 8):
    ws1.cell(row=tot_row, column=c).border = thin_border

# Auto column widths
for col in ws1.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.value:
            max_len = max(max_len, len(str(cell.value)))
    ws1.column_dimensions[col_letter].width = max(max_len + 3, 10)

ws1.column_dimensions["A"].width = 6
ws1.column_dimensions["C"].width = 38
ws1.column_dimensions["G"].width = 24

# Sheet 2: Ringkasan Pajak & Operasional
ws2 = wb.create_sheet(title="Ringkasan Biaya")
ws2.views.sheetView[0].showGridLines = True
ws2["A1"] = "RINGKASAN BIAYA OPERASIONAL"
ws2["A1"].font = Font(size=13, bold=True)
ws2.append(["Item Biaya", "Alokasi", "Realisasi", "Persentase"])
ws2.append(["Gaji & Tunjangan", 250000000, 245000000, "98%"])
ws2.append(["Operasional Server & Cloud", 85000000, 81000000, "95%"])
ws2.append(["Pemasaran & Kampanye Digital", 45000000, 43500000, "96%"])

wb.save("contoh_laporan_keuangan.xlsx")
print("File 'contoh_laporan_keuangan.xlsx' berhasil dibuat!")
